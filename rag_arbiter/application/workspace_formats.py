"""Controlled formats; edits operate on an immutable workspace revision only."""
import io, zipfile,re,copy
from pathlib import Path
from lxml import etree
from ..documents import digest

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS={'w':W}

def docx_xml(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist())>100*1024*1024:raise ValueError('DOCUMENT_TOO_LARGE')
        if any(i.filename.endswith('vbaProject.bin') for i in z.infolist()):raise ValueError('MACROS_NOT_SUPPORTED')
        return etree.fromstring(z.read('word/document.xml'),etree.XMLParser(resolve_entities=False,no_network=True))

def blocks(data,fmt):
    if fmt=='docx':
        root=docx_xml(data);body=root.find('w:body',NS);rows=[]
        for element in body:
            if element.tag not in ('{'+W+'}p','{'+W+'}tbl'):continue
            text=' '.join(element.xpath('.//w:t/text()',namespaces=NS))
            rows.append(dict(block_id=digest([len(rows),text]),reading_order=len(rows),block_type='table' if element.tag.endswith('tbl') else 'text',text=text,section=''))
        return {'blocks':rows}
    if fmt in ('txt','md'):
        return {'blocks':[dict(block_id=digest([i,t]),reading_order=i,block_type='text',text=t,section='') for i,t in enumerate(data.decode('utf-8').splitlines())]}
    if fmt=='pdf':
        import pypdfium2 as pdfium
        doc=pdfium.PdfDocument(data);texts=[]
        try:
            for i in range(len(doc)):
                page=doc[i];tp=page.get_textpage()
                try:texts.append(tp.get_text_range())
                finally:tp.close();page.close()
        finally:doc.close()
        return blocks('\n'.join(texts).encode(),'txt')
    raise ValueError('UNSUPPORTED_FORMAT')

def structured_content(doc,text):
    """Render explicit headings/lists/tables, keeping plain prose as prose."""
    def rich(paragraph,line):
        for i,part in enumerate(re.split(r'\*\*(.+?)\*\*',line)):
            run=paragraph.add_run(part);run.bold=bool(i%2)
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i];i+=1
        if line.strip().startswith('|') and i<len(lines) and re.fullmatch(r'[| :\-]+',lines[i].strip()):
            rows=[[c.strip() for c in line.strip().strip('|').split('|')]];i+=1
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')]);i+=1
            table=doc.add_table(rows=len(rows),cols=max(map(len,rows)));table.style='Table Grid'
            for r,values in enumerate(rows):
                for c,value in enumerate(values):rich(table.cell(r,c).paragraphs[0],value)
            continue
        heading=re.match(r'^(#{1,6})\s+(.+)',line)
        bullet=re.match(r'^\s*[-*]\s+(.+)',line);number=re.match(r'^\s*\d+[.)]\s+(.+)',line)
        if heading:rich(doc.add_paragraph(style='Heading '+str(min(len(heading[1]),3))),heading[2])
        elif bullet or number:rich(doc.add_paragraph(style='List Bullet' if bullet else 'List Number'),(bullet or number)[1])
        elif line.strip() not in ('---','***'):rich(doc.add_paragraph(),line)


def create(fmt,title,text):
    if fmt in ('txt','md'):return (title+'\n\n'+text).encode('utf-8')
    out=io.BytesIO()
    if fmt=='docx':
        from docx import Document
        doc=Document();doc.add_heading(title,0)
        structured_content(doc,text)
        doc.save(out)
    elif fmt=='pdf':
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        import reportlab,textwrap
        font=Path('C:/Windows/Fonts/arial.ttf')
        if not font.exists():font=Path(reportlab.__file__).parent/'fonts'/'Vera.ttf'
        if 'WorkspaceFont' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('WorkspaceFont',str(font)))
        c=canvas.Canvas(out);c.setFont('WorkspaceFont',11);y=800
        for line in (title+'\n\n'+text).splitlines():
            for part in textwrap.wrap(line,85) or ['']:
                if y<50:c.showPage();c.setFont('WorkspaceFont',11);y=800
                c.drawString(40,y,part);y-=16
        c.save()
    else:raise ValueError('UNSUPPORTED_FORMAT')
    return out.getvalue()

def edit(data,fmt,replacements,append):
    if fmt not in ('docx','txt','md'):raise ValueError('SAFE_EDIT_UNSUPPORTED')
    if fmt!='docx':
        text=data.decode('utf-8')
        for old,new in replacements.items():
            if text.count(old)!=1:raise ValueError('AMBIGUOUS_EDIT_TARGET')
            text=text.replace(old,new,1)
        return (text+('\n'+append if append else '')).encode()
    root=docx_xml(data)
    for old,new in replacements.items():
        matches=[p for p in root.findall('.//w:p',NS) if ''.join(p.xpath('.//w:t/text()',namespaces=NS))==old]
        if len(matches)!=1:raise ValueError('AMBIGUOUS_EDIT_TARGET')
        nodes=matches[0].findall('.//w:t',NS)
        if not nodes:raise ValueError('EMPTY_EDIT_TARGET')
        # Targeted text only; all other package parts, runs/properties remain.
        nodes[0].text=new;nodes[0].set('{http://www.w3.org/XML/1998/namespace}space','preserve')
        for node in nodes[1:]:node.text=''
    body=root.find('w:body',NS)
    from docx import Document
    additions=Document();structured_content(additions,append)
    for element in additions.element.body:
        if element.tag=='{'+W+'}sectPr':continue
        element=copy.deepcopy(element);section=body.find('w:sectPr',NS)
        if section is None:body.append(element)
        else:body.insert(body.index(section),element)
    out=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src,zipfile.ZipFile(out,'w') as dest:
        for info in src.infolist():dest.writestr(info,etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True) if info.filename=='word/document.xml' else src.read(info.filename))
    return out.getvalue()
