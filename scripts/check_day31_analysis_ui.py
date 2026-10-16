"""Real inference + MCP + browser on synthetic local DOCX; not SMB E2E."""
import argparse,json,os,sys,time,threading
from pathlib import Path
from uuid import uuid4
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import uvicorn
from docx import Document
from playwright.sync_api import sync_playwright,expect
from rag_arbiter.config import Config
from rag_arbiter.web.app import create_app
from rag_arbiter.application.grounding import Citation,CitationBuilder

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--reuse');args=parser.parse_args()
    root=Path(args.reuse).resolve() if args.reuse else (Path('data/day31-e2e')/('analysis-real-'+uuid4().hex[:10])).resolve();root.mkdir(parents=True,exist_ok=True)
    share=root/'share';share.mkdir(exist_ok=True)
    for name,days in [('alpha',7),('beta',9),('outside',99)]:
        if args.reuse and (share/(name+'.docx')).exists():continue
        d=Document();d.add_heading('Договор '+name,0);d.add_heading('Уведомление',1)
        d.add_paragraph(f'По договору {name} сторона обязана уведомить другую сторону за {days} календарных дней.');d.save(share/(name+'.docx'))
    cfg=Config.load().model_copy(deep=True);cfg.device='cpu';cfg.reranker.device='cpu'
    for key,value in dict(sqlite_path=root/'metadata.db',qdrant_path=root/'qdrant',corpus_path=root/'corpus',web_data_path=root/'web',uploads_path=root/'uploads',manifest_path=root/'manifest.json',report_path=root/'report.html').items():setattr(cfg,key,value)
    cfg.network_sources.allowed_roots=[str(share)];cfg.network_sources.enabled=False;cfg.network_sources.stability_seconds=0
    os.environ['HF_HUB_OFFLINE']='1';os.environ['TRANSFORMERS_OFFLINE']='1'
    app=create_app(cfg);server=uvicorn.Server(uvicorn.Config(app,host='127.0.0.1',port=18766,log_level='error'))
    thread=threading.Thread(target=server.run,daemon=True);thread.start()
    evidence=dict(classification='LOCAL SYNTHETIC DOCX / REAL MCP / REAL EMBEDDINGS AND GENERATION',status='STARTED')
    def save(stage):
        evidence['status']=stage;(root/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8');print(stage,flush=True)
    try:
        deadline=time.monotonic()+30
        while not server.started and time.monotonic()<deadline:time.sleep(.1)
        svc=app.state.network;source=svc.page('network_sources')[0] if args.reuse else svc.add(dict(name='Synthetic analysis E2E',root_path=str(share)))
        scanid=svc.start(source['id']);deadline=time.monotonic()+600
        while time.monotonic()<deadline:
            with app.state.runs.db() as store:scan=store.get('network_scans',scanid)
            if scan['status'] not in ('QUEUED','RUNNING'):break
            time.sleep(.5)
        assert scan['status']=='COMPLETED',scan
        rows=svc.page('network_files',source['id']);ids={r['filename']:r['document_id'] for r in rows};chosen={ids['alpha.docx'],ids['beta.docx']};save('INDEXED')
        sid=app.state.chat.create(None).session_id
        errors=[]
        with sync_playwright() as p:
            browser=p.chromium.launch(channel='msedge',headless=True);page=browser.new_page(viewport={'width':1440,'height':900});page.on('pageerror',lambda e:errors.append(str(e)))
            page.goto('http://127.0.0.1:18766/#sources');expect(page.locator('#sources-content')).to_contain_text('Проиндексировано: 3',timeout=15000)
            page.goto('http://127.0.0.1:18766/chat?session_id='+sid)
            assert page.locator('a[href="/catalog"]').count()==0
            page.locator('#choose-documents').click();expect(page.locator('[data-document-id]')).to_have_count(3)
            for did in chosen:page.locator('[data-document-id="'+did+'"]').check()
            page.locator('#selector-apply').click();expect(page.locator('#document-selector')).not_to_be_visible()
            page.reload();expect(page.locator('#chat-document-scope')).to_contain_text('2 документов')
            page.locator('#chat-input').fill('За сколько календарных дней сторона обязана уведомить другую сторону?');page.locator('#chat-submit').click()
            deadline=time.monotonic()+900;result=None
            while time.monotonic()<deadline:
                turns=app.state.chat.load(sid)['turns']
                if turns and turns[-1]['status'] not in ('QUEUED','RUNNING'):result=turns[-1];break
                page.wait_for_timeout(500)
            assert result,'generation timeout'
            evidence['turn']=result;save('ANSWER_'+result['grounding_status'])
            assert result['grounding_status']=='GROUNDED',result.get('response')
            assert set(result['trace']['eligible_document_ids'])==chosen
            assert {r['document_id'] for r in result['trace']['retrieved_candidates']}<=chosen
            assert result['citations_json']
            with app.state.runs.db() as store:assert all(CitationBuilder(store).exact(Citation(**c)) for c in result['citations_json'])
            expect(page.locator('#chat-submit')).to_be_enabled(timeout=10000)
            page.locator('summary').filter(has_text='Источники ·').click()
            page.locator('.chat-source').first.click();expect(page.locator('#chat-viewer')).to_be_visible();page.locator('#close-viewer').click()
            page.screenshot(path=str(root/'chat.png'))
            page.locator('#clear-chat-document').click();expect(page.locator('#chat-document-scope')).to_be_hidden()
            current=app.state.chat.load(sid)['session'];assert not current['selected_document_ids']
            ready=page.request.get(f'http://127.0.0.1:18766/api/chat/sessions/{sid}/documents?all_ready=true').json()['document_ids'];assert len(ready)==3
            for w,h in [(1366,768),(390,844)]:
                page.set_viewport_size(dict(width=w,height=h));assert page.locator('#chat-send').bounding_box()['y']<h
            assert not errors,errors;browser.close()
        evidence['selected_document_ids']=sorted(chosen);evidence['restored_document_count']=len(ready);save('PASS')
    except BaseException as e:evidence['failure']=str(e);save('FAILED');raise
    finally:server.should_exit=True;thread.join(timeout=60)
    print(root/'evidence.json',flush=True)
if __name__=='__main__':main()
