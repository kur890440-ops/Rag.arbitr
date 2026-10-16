"""Opt-in real SMB E2E. Writes one uniquely named synthetic DOCX in --source only.
Uses isolated local SQLite/Qdrant, existing app/SyncService/Chat, real inference. Only synthetic content can reach the configured cloud provider.
"""
import argparse
import json
import os
import sys
import time
from pathlib import Path
from uuid import uuid4
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient
from docx import Document
from rag_arbiter.config import Config
from rag_arbiter.web.app import create_app
from rag_arbiter.llm import MiniMaxLLMProvider
from rag_arbiter.application.grounding import Citation, CitationBuilder
from rag_arbiter.network.drive_paths import resolve_drive_path


def main():
    args=argparse.ArgumentParser();args.add_argument('--source',required=True);parsed=args.parse_args()
    root=Path(resolve_drive_path(parsed.source))
    assert root.is_dir(), 'Authorized test folder must exist'
    folder=Path('data/day31-e2e')/('assistant-real-'+uuid4().hex[:10]);folder.mkdir(parents=True)
    folder=folder.resolve()
    cfg=Config.load().model_copy(deep=True)
    cfg.device='cpu';cfg.reranker.device='cpu'
    cfg.sqlite_path=folder/'metadata.db';cfg.qdrant_path=folder/'qdrant';cfg.corpus_path=folder/'corpus'
    cfg.web_data_path=folder/'web';cfg.uploads_path=folder/'uploads';cfg.manifest_path=folder/'manifest.json';cfg.report_path=folder/'report.html'
    cfg.network_sources.allowed_roots=[str(root)];cfg.network_sources.enabled=True;cfg.network_sources.stability_seconds=1
    filename='day31-assistant-e2e-'+uuid4().hex[:10]+'.docx'
    fact='Срок уведомления по тестовому договору составляет семь календарных дней.'
    question='Какой срок уведомления установлен по тестовому договору?'
    evidence={'classification':'REAL SMB', 'generation':'configured MiniMax; synthetic test text only', 'test_filename':filename,'question':question,'fact':fact}
    output=folder/'evidence.json'
    def save(stage):
        evidence['stage']=stage;output.write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8');print(stage,flush=True)
    def wait_scan(service,sid=None):
        deadline=time.monotonic()+400
        while time.monotonic()<deadline:
            rows=service.page('network_scans',source['id'],limit=20)
            row=next((r for r in rows if r['scan_id']==sid),None) if sid else next((r for r in rows if r['ingestion_completed']),None)
            if row and row['status'] not in ('QUEUED','RUNNING'):return row
            time.sleep(.3)
        raise TimeoutError('SMB scan/index timeout')
    os.environ['HF_HUB_OFFLINE']='1';os.environ['TRANSFORMERS_OFFLINE']='1'
    with TestClient(create_app(cfg,llm_factory=lambda c:MiniMaxLLMProvider(c)),base_url='http://127.0.0.1',headers={'X-RAG-Request':'1'},client=('127.0.0.1',50000)) as client:
        sync=client.app.state.network;chat=client.app.state.chat
        try:
            access=client.post('/api/sources/check-path',json={'root_path':parsed.source});assert access.status_code==200,access.text
            evidence['access_status']=access.json()['status'];save('ACCESS_OK')
            source=sync.add(dict(name='Day31 synthetic SMB Chat test',root_path=str(root),include_patterns=[filename],scan_interval=5))
            evidence['source_id']=source['id']
            first=wait_scan(sync,sync.start(source['id']));evidence['initial_scan']=first
            assert first['files_discovered']==0
            doc=Document();doc.add_heading('Тестовый договор Day31',0);doc.add_heading('Срок уведомления',1);doc.add_paragraph(fact)
            path=root/filename
            assert not path.exists()
            with path.open('xb') as stream:doc.save(stream)
            save('SYNTHETIC_DOCX_CREATED')
            indexed=wait_scan(sync);evidence['automatic_scan']=indexed
            cfg.network_sources.enabled=False
            row=sync.page('network_files',source['id'])[0];evidence['file']={k:row.get(k) for k in ('document_id','run_id','status','indexing_status','chunk_count','format')}
            assert indexed['new_count']==indexed['ingestion_completed']==1,indexed
            assert row['indexing_status']=='SUCCESS';save('INDEXED')
            sid=chat.create(None).session_id
            chat.select_source(sid,'NETWORK',source['id'],0);evidence['session_id']=sid
            panel=client.get(f'/ui/chat/{sid}/discovery');assert filename in panel.text and 'INDEXED' in panel.text
            evidence['chat_panel_matches_registry']=True
            turn=chat.start(sid,'/help '+question,0);chat.futures[turn.turn_id].result(timeout=900)
            result=chat.load(sid)['turns'][-1]
            evidence['chat_turn']=result
            evidence['answer']=result.get('response',{}).get('answer_text')
            evidence['citations']=result['citations_json']
            with sync.runs.db() as store:
                evidence['citations_exact']=bool(result['citations_json']) and all(CitationBuilder(store).exact(Citation(**c)) for c in result['citations_json'])
            save('CHAT_'+result['grounding_status'])
            assert result['grounding_status']=='GROUNDED',result.get('trace',{}).get('error',result.get('response'))
            assert evidence['citations_exact']
            assert {s['document_id'] for s in result['sources_json']}=={row['document_id']}
            before_scans={r['scan_id'] for r in sync.page('network_scans',source['id'])}
            status_question='/help Какие документы обнаружены при последнем scan?'
            state=chat.load(sid)['task_state']
            status_turn=chat.start(sid,status_question,state['version'])
            chat.futures[status_turn.turn_id].result(timeout=180)
            status=chat.load(sid)['turns'][-1]
            evidence['archive_status_turn']=status
            assert status['response']['status']=='ARCHIVE_STATUS',status['response']
            assert filename in status['response']['answer_text']
            assert status['trace']['archive_status']['counts']['indexed']==1
            assert status['trace']['archive_status']['files'][0]['last_scan']
            assert status['comparison_run_id'] is None
            assert before_scans=={r['scan_id'] for r in sync.page('network_scans',source['id'])}
            with sync.runs.db() as store:
                canonical=store.get('documents',row['document_id'])
                assert canonical and canonical['blocks']
                evidence['canonical_blocks']=len(canonical['blocks'])
                evidence['canonical_fact_present']=any(fact in b['text'] for b in canonical['blocks'])
                assert evidence['canonical_fact_present']
            save('STATUS_VERIFIED')
            repeat=wait_scan(sync,sync.start(source['id']));evidence['repeat_scan']=repeat
            assert repeat['unchanged_count']==1 and repeat['ingestion_queued']==0;save('UNCHANGED')
            save('PASS')
        except BaseException as exc:
            evidence['failure_type']=type(exc).__name__;save('FAILED');raise
    print('Evidence:',output,flush=True)

if __name__=='__main__':main()
