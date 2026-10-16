"""REAL SMB/MCP/Docling/BGE/local Qwen/Edge; synthetic authorized content only."""
import json,os,sys,time,threading,argparse
from pathlib import Path
from uuid import uuid4
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import uvicorn
from docx import Document
from playwright.sync_api import sync_playwright,expect
from rag_arbiter.config import Config
from rag_arbiter.web.app import create_app
from rag_arbiter.application.document_review import DocumentReviewService
from rag_arbiter.application.grounding import Citation,CitationBuilder
from rag_arbiter.application.views import ResultsService

def make(path,v):
    doc=Document();doc.add_heading('Агентский договор — синтетический тест Day32',0)
    doc.add_paragraph('1. Агент обязан направлять отчёт ежемесячно.')
    doc.add_paragraph('2. Передать оборудование на склад.' if v=='A' else '2. Разрешить использование товарного знака.')
    table=doc.add_table(rows=2,cols=2)
    for row,values in zip(table.rows,[['Условие','Значение'],['Агентское вознаграждение','5%' if v=='A' else '3%']]):
        for cell,value in zip(row.cells,values):cell.text=value
    doc.save(path)

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--source',required=True);parser.add_argument('--model');args=parser.parse_args()
    from rag_arbiter.network.drive_paths import resolve_drive_path
    authorized=Path(resolve_drive_path(args.source));assert authorized.is_dir()
    key=uuid4().hex[:10];share=authorized/('day32-review-'+key);share.mkdir()
    folder=(Path('data/day32-e2e')/key).resolve();folder.mkdir(parents=True)
    cfg=Config.load().model_copy(deep=True);cfg.device='cpu';cfg.reranker.device='cpu'
    for name,value in dict(sqlite_path='metadata.db',qdrant_path='qdrant',corpus_path='corpus',web_data_path='web',uploads_path='uploads',manifest_path='manifest.json',snapshot_path='snapshot.json',report_path='report.html').items():setattr(cfg,name,folder/value)
    if args.model:cfg.llm.local.model=args.model
    cfg.network_sources.allowed_roots=[str(share)];cfg.network_sources.enabled=True;cfg.network_sources.stability_seconds=.3
    os.environ['HF_HUB_OFFLINE']='1';os.environ['TRANSFORMERS_OFFLINE']='1'
    app=create_app(cfg);server=uvicorn.Server(uvicorn.Config(app,host='127.0.0.1',port=18770,log_level='error'))
    thread=threading.Thread(target=server.run,daemon=True);thread.start()
    evidence=dict(classification='REAL SMB / MCP / DOCLING / BGE / LOCAL QWEN / EDGE',cloud_used=False,test_folder=str(share),model=cfg.llm.local.model)
    def save(stage):
        evidence['stage']=stage;(folder/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8');print(stage,flush=True)
    def wait_job(n):
        deadline=time.monotonic()+600
        while time.monotonic()<deadline:
            with app.state.runs.db() as store:jobs=store.all('document_reviews')
            if len(jobs)>=n and jobs[n-1]['status'] in ('COMPLETED','FAILED'):return jobs[n-1]
            time.sleep(1)
        raise AssertionError('Review timeout')
    try:
        for _ in range(60):
            if server.started:break
            time.sleep(.5)
        assert server.started
        net=app.state.network;chat=app.state.chat
        # Record real local provider calls, without changing generation or return values.
        original=net.reviews.provider_factory;calls=[]
        def provider(c):
            impl=original(c)
            class Recorded:
                def generate(self,r):
                    response=impl.generate(r);calls.append(response.model_dump());(folder/'local-calls.json').write_text(json.dumps(calls,ensure_ascii=False,indent=2),encoding='utf-8');return response
            return Recorded()
        net.reviews.provider_factory=provider
        source=net.add(dict(name='Day32 SMB review',root_path=str(share),scan_interval=3))
        sid=chat.create(None).session_id;other=chat.create(None).session_id;chat.select_source(sid,'NETWORK',source['id'],0)
        with sync_playwright() as p:
            browser=p.chromium.launch(channel='msedge',headless=True);page=browser.new_page(viewport={'width':1440,'height':900});errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
            page.goto('http://127.0.0.1:18770/chat?session_id='+sid)
            path=share/'Агентский договор Day32.docx';make(path,'A');save('A_CREATED_AUTO_SCAN')
            first=wait_job(1);evidence['new_review']=first;save('A_REVIEW_'+first['status']);assert first['status']=='COMPLETED',first.get('error')
            assert any('обязан' in x['text'] or 'оборудование' in x['text'] or '5%' in x['text'] for x in first['result']['summary'])
            expect(page.locator('#document-review-feed')).to_contain_text('Краткое содержание',timeout=12000)
            page.screenshot(path=str(folder/'new-summary.png'))
            row=net.page('network_files',source['id'])[0];did=row['document_id'];chat.select_source(sid,'NETWORK',source['id'],1,[did])
            page.reload();other_before=chat.load(other);memory=chat.load(sid)['task_state']
            with net.runs.db() as store:
                old=store.get('document_version_contents',first['new_version_id']);old_copy=json.dumps(old,sort_keys=True)
                block=next(b for b in old['blocks'] if b['block_type']=='table')
                citation=Citation(citation_id='e2e-old',claim_id='C1',source_id='S1',document_id=did,chunk_id='',anchor_chunk_ids=[],file_title=path.name,section=block['section'],page_start=0,page_end=0,quote=block['text'],locator_kind='structure',locator=dict(block_id=block['block_id'],recognition_id=block['recognition_id'],start=0,end=len(block['text'])))
                assert CitationBuilder(store).exact(citation)
            make(path,'B');save('B_CREATED_AUTO_SCAN')
            second=wait_job(2);evidence['modified_review']=second;save('B_REVIEW_'+second['status']);assert second['status']=='COMPLETED',second.get('error')
            assert second['counts']==dict(ADDED=1,REMOVED=1,MODIFIED=1,MOVED=0),second['changes']
            assert any(c['old_value']=='5%' and c['new_value']=='3%' for change in second['changes'] for c in change['table_changes'])
            expect(page.locator('#document-review-feed')).to_contain_text('Изменён документ',timeout=12000)
            expect(page.locator('#discovery-files')).to_contain_text('Новая версия',timeout=12000)
            assert chat.load(sid)['session']['selected_document_ids']==[did]
            with net.runs.db() as store:
                assert json.dumps(store.get('document_version_contents',first['new_version_id']),sort_keys=True)==old_copy
                assert CitationBuilder(store).exact(citation)
            newdid=net.page('network_files',source['id'])[0]['document_id']
            pinned=ResultsService(net.runs).chat_snapshot(None,source['id'],source_scope='NETWORK',selected_document_ids=[did])
            current=ResultsService(net.runs).chat_snapshot(None,source['id'],source_scope='NETWORK')
            assert set(pinned['document_runs'])=={did} and set(current['document_runs'])=={newdid}
            details=page.locator('#document-review-feed details').first;details.locator('summary').click();page.screenshot(path=str(folder/'modified-diff.png'))
            details.locator('.chat-source').first.click();expect(page.locator('#chat-viewer')).to_be_visible();page.screenshot(path=str(folder/'old-source.png'));page.locator('#close-viewer').click()
            page.locator('#discovery-add').click();expect(page.locator('#chat-discovery-notice')).to_be_hidden(timeout=15000)
            assert chat.load(sid)['session']['selected_document_ids']==[newdid] and chat.load(sid)['task_state']==memory and chat.load(other)==other_before
            prior_calls=len(calls);time.sleep(12)
            with net.runs.db() as store:assert len(store.all('document_reviews'))==2
            assert len(calls)==prior_calls
            deadline=time.monotonic()+60
            while time.monotonic()<deadline:
                completed=[s for s in net.page('network_scans',source['id'],limit=20) if s['status']=='COMPLETED' and s['started_at']>second['finished_at']]
                if completed and completed[0]['unchanged_count']==1:break
                time.sleep(1)
            assert completed and completed[0]['unchanged_count']==1
            latest=completed[0]
            page.set_viewport_size({'width':390,'height':844});page.screenshot(path=str(folder/'mobile.png'));assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
            assert not errors,errors
            evidence.update(session_id=sid,source_id=source['id'],old_version_preserved=True,old_citation_exact=True,versions_not_mixed=True,selection_updated_only_on_confirmation=True,other_session_unchanged=True,automatic_chat_display=True,unchanged_no_review=True,provider_calls=len(calls),last_scan=latest)
            browser.close();save('PASS')
    except BaseException as exc:evidence['error']=str(exc);save('FAILED');raise
    finally:cfg.network_sources.enabled=False;server.should_exit=True;thread.join(30)
    print(folder,flush=True)
if __name__=='__main__':main()
