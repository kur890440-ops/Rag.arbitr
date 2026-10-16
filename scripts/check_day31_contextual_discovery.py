"""Real SMB + MCP + local BGE/Qwen + Edge, isolated metadata and test documents."""
import json,os,sys,time,threading
from pathlib import Path
from uuid import uuid4
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
import uvicorn
from playwright.sync_api import sync_playwright,expect
from rag_arbiter.config import Config
from rag_arbiter.web.app import create_app
from rag_arbiter.local_llm import LocalLLMProvider
from rag_arbiter.application.grounding import Citation,CitationBuilder


def main():
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('--source',required=True);parser.add_argument('--resume');parser.add_argument('--reuse-initial');parser.add_argument('--model');parser.add_argument('--view-result',action='store_true');args=parser.parse_args()
    reuse=args.resume or args.reuse_initial
    from rag_arbiter.network.drive_paths import resolve_drive_path
    authorized=Path(resolve_drive_path(args.source));assert authorized.is_dir()
    key=uuid4().hex[:10];share=authorized/('day31-context-'+key)
    if reuse:
        previous=json.loads(Path(reuse).read_text(encoding='utf-8'));share=Path(previous['test_folder'])
    else:share.mkdir()
    assert share.resolve().is_relative_to(authorized.resolve())
    folder=(Path('data/day31-e2e')/('context-real-'+key)).resolve();folder.mkdir(parents=True)
    if reuse:folder=Path(reuse).resolve().parent
    cfg=Config.load().model_copy(deep=True)
    if args.model:cfg.llm.local.model=args.model
    cfg.device='cpu';cfg.reranker.device='cpu'
    for name,path in dict(sqlite_path='metadata.db',qdrant_path='qdrant',corpus_path='corpus',web_data_path='web',uploads_path='uploads',manifest_path='manifest.json',snapshot_path='snapshot.json',report_path='report.html').items():setattr(cfg,name,folder/path)
    cfg.network_sources.allowed_roots=[str(share)];cfg.network_sources.enabled=False;cfg.network_sources.stability_seconds=.2
    os.environ['HF_HUB_OFFLINE']='1';os.environ['TRANSFORMERS_OFFLINE']='1'
    for i in range(0 if reuse else 16):
        doc=Document();doc.add_heading('Фоновая запись '+str(i+1),0);doc.add_paragraph('Место хранения оборудования — северный склад. Номер записи '+str(i+1)+'.');doc.save(share/f'base-{i+1:02}.docx')
    calls=[]
    def provider(c):
        implementation=LocalLLMProvider(c.local)
        class Recorded:
            def generate(self,request):
                result=implementation.generate(request)
                calls.append(dict(purpose=request.question[:60],result=result.model_dump()))
                (folder/'local-calls.json').write_text(json.dumps(calls,ensure_ascii=False,indent=2),encoding='utf-8')
                return result
        return Recorded()
    app=create_app(cfg,llm_factory=provider)
    server=uvicorn.Server(uvicorn.Config(app,host='127.0.0.1',port=18769,log_level='error'))
    thread=threading.Thread(target=server.run,daemon=True);thread.start()
    evidence={'classification':'REAL SMB / MCP / BGE / LOCAL QWEN / EDGE','test_folder':str(share),'cloud_used':False,'local_model':cfg.llm.local.model}
    def save(stage):
        evidence['stage']=stage;(folder/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8');print(stage,flush=True)
    try:
        deadline=time.monotonic()+30
        while not server.started and time.monotonic()<deadline:time.sleep(.1)
        assert server.started
        net=app.state.network;chat=app.state.chat
        if args.resume:
            source=net.source(previous['source_id']);sid=previous['session_id'];evidence.update(previous)
        else:
            source=net.page('network_sources')[0] if args.reuse_initial else net.add(dict(name='SMB contextual E2E',root_path=str(share),scan_interval=3))
            scan_id=net.start(source['id']);save('INITIAL_16_INDEXING')
            deadline=time.monotonic()+900;count=-1
            while time.monotonic()<deadline:
                latest=net.page('network_scans',source['id'],limit=1)[0]
                if latest['ingestion_completed']!=count:count=latest['ingestion_completed'];print('indexed',count,flush=True)
                if latest['status'] not in ('RUNNING','QUEUED'):break
                time.sleep(1)
            assert latest['ingestion_completed']+latest['unchanged_count']==16 and latest['status']=='COMPLETED',latest
            ids=[r['document_id'] for r in net.page('network_files',source['id'])]
            sid=chat.create(None).session_id;other=chat.create(None).session_id
            chat.select_source(sid,'NETWORK',source['id'],0,ids)
            other_before=chat.load(other);memory_before=chat.load(sid)['task_state']
            evidence.update(source_id=source['id'],session_id=sid,initial_selection_count=len(ids),initial_scan=latest)
        with sync_playwright() as p:
            browser=p.chromium.launch(channel='msedge',headless=True);page=browser.new_page(viewport={'width':1440,'height':900});errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
            page.goto('http://127.0.0.1:18769/chat?session_id='+sid)
            if not args.resume:
                page.evaluate("window.toastSeen=false;new MutationObserver(()=>{if(!document.querySelector('#knowledge-toast').hidden)window.toastSeen=true}).observe(document.querySelector('#knowledge-toast'),{attributes:true})")
                expect(page.locator('#chat-document-scope')).to_contain_text('16 документов')
                expect(page.locator('#chat-discovery-notice')).to_be_hidden()
                cfg.network_sources.enabled=True
                filename='Контур-17.docx';doc=Document();doc.add_heading('Тестовый договор Контур-17',0);doc.add_paragraph('Срок уведомления по договору Контур-17 составляет семь календарных дней.');doc.save(share/filename)
                save('NEW_DOCX_CREATED_AUTOMATIC_SCAN_PENDING')
                expect(page.locator('#chat-discovery-notice')).to_contain_text(filename,timeout=120000)
                evidence['appeared_without_click']=True
                assert chat.load(sid)['session']['selected_document_ids']==ids
                page.screenshot(path=str(folder/'processing.png'))
                expect(page.locator('#discovery-add')).to_be_enabled(timeout=240000)
                expect(page.locator('#discovery-files')).to_contain_text('INDEXED')
                page.screenshot(path=str(folder/'ready.png'))
                page.set_viewport_size({'width':390,'height':844})
                assert page.locator('#chat-input').bounding_box()['height'] >= 40
                assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
                page.screenshot(path=str(folder/'ready-mobile.png'))
                page.set_viewport_size({'width':1440,'height':900})
                assert not page.evaluate('window.toastSeen')
                page.reload();expect(page.locator('#chat-discovery-notice')).to_contain_text(filename,timeout=10000)
                page.locator('#discovery-add').click()
                expect(page.locator('#chat-document-scope')).to_contain_text('17 документов',timeout=10000)
                expect(page.locator('#chat-discovery-notice')).to_be_hidden()
                assert chat.load(sid)['task_state']==memory_before
                assert chat.load(other)==other_before
                page.reload();expect(page.locator('#chat-document-scope')).to_contain_text('17 документов')
                page.wait_for_timeout(3500);expect(page.locator('#chat-discovery-notice')).to_be_hidden();expect(page.locator('#knowledge-toast')).to_be_hidden()
                evidence.update(final_selection_count=17,memory_unchanged=True,other_session_unchanged=True,no_duplicate_toast=True,reload_persisted=True)
                save('SELECTION_17_CONFIRMED')
            filename='Контур-17.docx'
            question='Сколько календарных дней составляет срок уведомления по договору Контур-17?' 
            if not args.view_result:
                page.locator('#chat-input').fill(question);page.locator('#chat-input').press('Enter')
                deadline=time.monotonic()+600
                while time.monotonic()<deadline:
                    page.wait_for_timeout(1000);turns=chat.load(sid)['turns']
                    if turns and turns[-1]['status'] not in ('QUEUED','RUNNING'):break
            turn=chat.load(sid)['turns'][-1];evidence['turn']=turn
            with net.runs.db() as store:
                exact=bool(turn['citations_json']) and all(CitationBuilder(store).exact(Citation(**c)) for c in turn['citations_json'])
            newrow=next(r for r in net.page('network_files',source['id']) if r['filename']==filename)
            evidence['new_document_id']=newrow['document_id'];evidence['citations_exact']=exact
            save('ANSWER_'+turn['grounding_status'])
            assert turn['grounding_status']=='GROUNDED' and exact,turn.get('response')
            assert newrow['document_id'] in {r['document_id'] for r in turn['sources_json']}
            expect(page.locator('.chat-message.assistant').last).to_contain_text(turn['response']['answer_text'].split(' [S')[0],timeout=15000)
            expect(page.locator('#chat-submit')).to_be_enabled(timeout=15000)
            page.screenshot(path=str(folder/'answer.png'));assert not errors,errors
            browser.close()
        evidence.pop('error_type',None)
        save('PASS')
    except BaseException as exc:
        evidence['error_type']=type(exc).__name__;save('FAILED');raise
    finally:
        cfg.network_sources.enabled=False;server.should_exit=True;thread.join(30)
    print(folder,flush=True)

if __name__=='__main__':main()

