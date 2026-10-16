import copy,json,time
import pytest
from docx import Document
from rag_arbiter.application.document_diff import structured_diff,table_grid
from rag_arbiter.application.document_review import FIELDS,REVIEW_FIELDS
from rag_arbiter.application.views import ResultsService
from rag_arbiter.application.contextual_discovery import ContextualDiscovery
from rag_arbiter.application.grounding import CitationBuilder,Citation
from rag_arbiter.llm import LLMResult
from test_rag import ragweb
from test_network_sources import scan


def blocks(texts):
    return {'blocks':[dict(block_id=str(i)+t,text=t,block_type='text',reading_order=i,section='Условия') for i,t in enumerate(texts)]}

def diff(a,b):return structured_diff('logical','A','B',blocks(a),blocks(b))

@pytest.mark.parametrize('a,b,types',[
    (['Сторона А.'],['Сторона А.','Возврат оборудования.'],['ADDED']),
    (['Сторона А.','Возврат оборудования.'],['Сторона А.'],['REMOVED']),
    (['Вознаграждение агента составляет 5%.'],['Вознаграждение агента составляет 3%.'],['MODIFIED']),
    (['Первый уникальный текст.','Второй самостоятельный пункт.','Третий договорный раздел.'],['Третий договорный раздел.','Первый уникальный текст.','Второй самостоятельный пункт.'],['MOVED']),
    (['1. Обязательство поставщика.','2. Обязанность покупателя.'],['4. Обязательство поставщика.','5. Обязанность покупателя.'],[]),
    (['5 процентов от суммы.'],['3 процента от суммы.'],['MODIFIED']),
    (['10.10.2026 срок оплаты.'],['11.10.2026 срок оплаты.'],['MODIFIED']),
    (['  Оплата   в срок.'],['Оплата в срок.'],[]),
])
def test_structural_cases(a,b,types):assert sorted(c['change_type'] for c in diff(a,b))==sorted(types)

def test_table_cells_and_ambiguous():
    a=blocks(['| Условие | Значение |\n|---|---|\n| Вознаграждение | 5% |'])
    b=blocks(['| Условие | Значение |\n|---|---|\n| Вознаграждение | 3% |'])
    for d in (a,b):d['blocks'][0]['block_type']='table'
    changes=structured_diff('d','a','b',a,b)
    assert changes[0]['table_changes']==[dict(row='Вознаграждение',column='Значение',old_value='5%',new_value='3%',change_type='MODIFIED')]
    changes=diff(['Повторный пункт.','Повторный пункт.'],['Повторный пункт.'])
    assert any(c['status']=='UNCERTAIN' for c in changes)

class FixtureReview:
    def __init__(self):self.calls=0;self.fail=False
    def generate(self,request):
        self.calls+=1
        if self.fail:return LLMResult(status='ERROR',model='MOCK')
        data=json.loads(request.question.split('\n',1)[1])
        result={k:[] for k in request.response_schema['properties']}
        result['summary']=[request.response_schema['properties']['summary']['items']['enum'][0]] if request.response_schema['properties']['summary']['items']['type']=='string' else [dict(text='Изменены условия.',source_references=['U1'])]
        return LLMResult(status='SUCCESS',model='MOCK',text=json.dumps(result,ensure_ascii=False))

def make(path,version):
    d=Document();d.add_heading('Агентский договор',0)
    d.add_paragraph('1. Агент обязан направлять отчёт ежемесячно.')
    if version=='A':d.add_paragraph('2. Передать оборудование на склад.')
    else:d.add_paragraph('2. Разрешить использование товарного знака.')
    d.add_paragraph('3. Вознаграждение агента составляет '+('5%' if version=='A' else '3%')+'.')
    t=d.add_table(rows=2,cols=2)
    for row,vals in zip(t.rows,[['Условие','Значение'],['Вознаграждение','5%' if version=='A' else '3%']]):
        for c,v in zip(row.cells,vals):c.text=v
    d.save(path)

def jobs(net,n):
    deadline=time.monotonic()+40
    while time.monotonic()<deadline:
        with net.runs.db() as store:rows=store.all('document_reviews')
        if len(rows)==n and all(r['status'] not in ('RUNNING','QUEUED') for r in rows):return rows
        time.sleep(.1)
    raise AssertionError(rows)

def test_version_review_scope_and_failure(ragweb,tmp_path):
    client,rid,_,_=ragweb;net=client.app.state.network;chat=client.app.state.chat
    provider=FixtureReview();net.reviews.provider_factory=lambda _:provider
    root=tmp_path/'review-share';root.mkdir();net.config.network_sources.allowed_roots.append(str(root));net.config.network_sources.stability_seconds=0
    source=net.add(dict(name='Review',root_path=str(root)));sid=chat.create(rid).session_id;other=chat.create(rid).session_id
    chat.select_source(sid,'NETWORK',source['id'],0)
    path=root/'Agreement.docx';make(path,'A');assert scan(net,source)['ingestion_completed']==1
    first=jobs(net,1)[0];assert first['status']=='COMPLETED',first
    assert first['result']['summary'] and first['source_references']
    row=net.page('network_files',source['id'])[0];did=row['document_id']
    chat.select_source(sid,'NETWORK',source['id'],1,[did]);before=chat.load(other)
    with net.runs.db() as store:
        old=store.get('document_version_contents',first['new_version_id']);saved=copy.deepcopy(old)
        b=next(b for b in old['blocks'] if '5%' in b['text'])
        citation=Citation(citation_id='test',claim_id='C1',source_id='S1',document_id=did,chunk_id='',anchor_chunk_ids=[],file_title='Agreement',section='',page_start=0,page_end=0,quote=b['text'],locator_kind='structure',locator=dict(block_id=b['block_id'],start=0,end=len(b['text']),recognition_id=b['recognition_id']))
        assert CitationBuilder(store).exact(citation)
    make(path,'B');assert scan(net,source)['ingestion_completed']==1
    changed=jobs(net,2)[1];assert changed['status']=='COMPLETED',changed
    assert changed['old_version_id']==first['new_version_id']
    assert changed['counts']['ADDED']==1 and changed['counts']['REMOVED']==1 and changed['counts']['MODIFIED']==2,changed['changes']
    with net.runs.db() as store:
        assert store.get('document_version_contents',first['new_version_id'])==saved
        assert CitationBuilder(store).exact(citation)
    newdid=net.page('network_files',source['id'])[0]['document_id']
    active=ResultsService(net.runs).chat_snapshot(rid,source['id'],source_scope='NETWORK')
    pinned=ResultsService(net.runs).chat_snapshot(rid,source['id'],source_scope='NETWORK',selected_document_ids=[did])
    assert set(active['document_runs'])=={newdid} and set(pinned['document_runs'])=={did}
    assert not set(active['indexes']['structure']['chunk_ids']) & set(pinned['indexes']['structure']['chunk_ids'])
    with pytest.raises(ValueError,match='одну версию'):
        chat.select_source(sid,'NETWORK',source['id'],2,[did,newdid])
    notice=ContextualDiscovery(chat).snapshot(sid);assert notice['files'][0]['change_type']=='MODIFIED'
    ContextualDiscovery(chat).decide(sid,'add',[f['event_id'] for f in notice['files']],notice['scope_version'],[newdid])
    assert chat.load(sid)['session']['selected_document_ids']==[newdid]
    assert chat.load(other)==before
    with net.runs.db() as store:
        store.put('document_version_contents',first['new_version_id'],{'tampered':True})
        assert store.get('document_version_contents',first['new_version_id'])==saved
    response=client.get('/api/chat/knowledge-updates',params={'after':0,'session_id':sid}).json()
    assert len(response['reviews'])==2
    calls=provider.calls;assert scan(net,source)['unchanged_count']==1;assert len(jobs(net,2))==2 and provider.calls==calls
    with net.runs.db() as store:prior=store.get('document_versions',first['new_version_id'])
    net.reviews.schedule(net.page('network_files',source['id'])[0],prior,modified=True)
    assert len(jobs(net,2))==2 and provider.calls==calls
    provider.fail=True
    make(root/'failed-review.docx','A');assert scan(net,source)['ingestion_completed']==1
    with net.runs.db() as store:count=len(store.all('document_reviews'))
    results=jobs(net,count);assert any(r['error']=='REVIEW_GENERATION_FAILED' for r in results)
    assert all(r['indexing_status']=='SUCCESS' for r in net.page('network_files',source['id']))
    net.config.network_sources.allowed_roots.remove(str(root));assert not net.reviews.feed(chat.load(sid)['session'])

def test_review_error_codes_and_bounded_coverage(config,tmp_path,monkeypatch):
    from contextlib import contextmanager
    from concurrent.futures import ThreadPoolExecutor
    from threading import RLock
    from rag_arbiter.storage import MetadataStore
    from rag_arbiter.application.document_review import DocumentReviewService,ReviewBatcher
    from rag_arbiter.documents import now
    root=tmp_path/'allowed';root.mkdir();config.network_sources.allowed_roots=[str(root)]
    class Runs:
        mutation_lock=RLock()
        executor=ThreadPoolExecutor(max_workers=1)
        @contextmanager
        def db(self):
            store=MetadataStore(config.sqlite_path)
            try:yield store
            finally:store.close()
    runs=Runs();runs.config=config;provider=FixtureReview();service=DocumentReviewService(runs,lambda _:provider)
    monkeypatch.setattr('rag_arbiter.application.document_review.ready',lambda *_:True)
    doc=dict(document_id='canonical',file_name='fixture.pdf',pages=[],page_count=0,locator_kind='structure',cache_identity='norm',parser_version='test',blocks=[dict(block_id='block',text='Условие оплаты. '*6000,section='Сроки',reading_order=0,page_number=0,block_type='text')])
    with runs.db() as store:
        store.put('network_sources','s',dict(id='s',root_path=str(root)))
        store.put('document_versions','v',dict(version_id='v',canonical_content_reference='v'))
        store.put('document_version_contents','v',doc)
    def execute(kind,old=None):
        key=str(time.monotonic());job=dict(review_id=key,document_id='d',source_id='s',document_name='fixture.pdf',old_version_id=old,new_version_id='v',kind=kind,status='QUEUED',created_at=now(),changes=[],result={})
        service.save(job);service.execute(key)
        with runs.db() as store:return store.get('document_reviews',key)
    missing=execute('MODIFIED');assert missing['error']=='VERSION_NOT_AVAILABLE'
    monkeypatch.setattr('rag_arbiter.application.document_review.ready',lambda *_:False)
    assert execute('NEW')['error']=='INDEX_NOT_READY'
    monkeypatch.setattr('rag_arbiter.application.document_review.ready',lambda *_:True)
    result=execute('NEW');assert result['status']=='COMPLETED',result
    assert result['coverage']['batches_completed']>1
    # Every exact input character is covered by the existing batcher; no truncation/RAG substitute.
    batcher=ReviewBatcher(config.llm,'NEW');batches=batcher.build([doc])
    assert ''.join(u.text for b in batches for u in b.units)==doc['blocks'][0]['text']
    huge={**doc,'blocks':[{**doc['blocks'][0],'text':'Short'}]*2001}
    with runs.db() as store:
        store.put('document_versions','huge',dict(version_id='huge',canonical_content_reference='huge'))
        store.put('document_version_contents','huge',huge)
        job=copy.deepcopy(result);job.update(review_id='large',new_version_id='huge',status='QUEUED');store.put('document_reviews','large',job)
    service.execute('large')
    with runs.db() as store:assert store.get('document_reviews','large')['error']=='DOCUMENT_TOO_LARGE'
    monkeypatch.setattr('rag_arbiter.application.document_review.structured_diff',lambda *_:(_ for _ in ()).throw(RuntimeError('test')))
    assert execute('MODIFIED','v')['error']=='DIFF_FAILED'
    runs.executor.shutdown()

def test_review_feed_respects_sources_and_selection(config,tmp_path):
    from contextlib import contextmanager
    from rag_arbiter.storage import MetadataStore
    from rag_arbiter.application.document_review import DocumentReviewService
    from rag_arbiter.documents import now
    class Runs:
        @contextmanager
        def db(self):
            store=MetadataStore(config.sqlite_path)
            try:yield store
            finally:store.close()
    runs=Runs();runs.config=config
    for sid in ('s1','s2'):
        root=tmp_path/sid;root.mkdir();config.network_sources.allowed_roots.append(str(root))
        with runs.db() as store:store.put('network_sources',sid,dict(id=sid,root_path=str(root)))
    with runs.db() as store:
        for n,sid in enumerate(('s1','s1','s2'),1):
            store.put('document_reviews',str(n),dict(review_id=str(n),document_id='d'+str(n),source_id=sid,created_at=now()))
            store.put('document_versions',str(n),dict(document_id='d'+str(n),canonical_document_id='c'+str(n)))
    service=object.__new__(DocumentReviewService);service.runs=runs
    session=dict(source_scope='NETWORK',source_id='s1',selected_document_ids=['c1'],created_at='2000')
    assert {r['review_id'] for r in service.feed(session)}=={'1','2'}  # Same-source NEW can be offered, not auto-selected.
    session.update(source_scope='ALL_DOCUMENTS',source_id=None)
    assert {r['review_id'] for r in service.feed(session)}=={'1'}
    session['selected_document_ids']=[]
    assert len(service.feed(session))==3
    session['source_scope']='MANUAL';assert not service.feed(session)
    session.update(source_scope='NETWORK',source_id='s2');assert {r['review_id'] for r in service.feed(session)}=={'3'}
