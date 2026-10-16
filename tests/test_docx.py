"""Real Docling/MCP/SQLite/Qdrant; inference doubles are explicitly MOCK TEST."""
import json
import re
from pathlib import Path
import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from rag_arbiter.ingestion import Ingestor
from rag_arbiter.storage import MetadataStore
from rag_arbiter.docx_adapter import MIME
from rag_arbiter.web.app import create_app
from rag_arbiter.llm import LLMResult
from rag_arbiter.application.chat_memory import TaskMemoryUpdater, ContextualQueryBuilder
from rag_arbiter.application.grounding import CitationBuilder, Citation
from test_web import factory, HEADERS
from test_recognition import FakeRuntime
from test_network_sources import scan
from test_chat import MemoryProvider
from dialogue_fixtures import install

FACT='Умышленное причинение ущерба не признаётся страховым случаем.'
QUESTION='Признаётся ли умышленное причинение ущерба страховым случаем?'


def make_docx(path, fact=FACT):
    doc=WordDocument()
    doc.add_heading('Правила страхования — тестовый документ',0)
    doc.add_heading('Исключения',1)
    doc.add_paragraph(fact)
    doc.add_paragraph('Уведомить страховщика.',style='List Bullet')
    doc.add_paragraph('Приложить документы.',style='List Number')
    table=doc.add_table(rows=2,cols=3)
    for row,values in zip(table.rows,[['Страховой случай','Исключение','Условие выплаты'],['Пожар','Умысел','Подтверждающие документы']]):
        for cell,value in zip(row.cells,values):cell.text=value
    doc.save(path)


class NoOCR:
    def recognize_page(self,*a):raise AssertionError('DOCX must not invoke OCR')
    def preflight(self):raise AssertionError('DOCX must not preflight Ollama')
    def close(self):pass


def test_real_docling_canonical_structure_table_cache(config):
    path=config.corpus_path/'insurance.docx';make_docx(path)
    store=MetadataStore(config.sqlite_path)
    try:
        docs,stats=Ingestor(config,store,NoOCR()).ingest([path])
        assert not stats['errors'],stats
        doc=docs[0]
        assert doc.page_count==0 and doc.pages==[] and doc.locator_kind=='structure'
        assert doc.content_type==MIME
        assert all(b.page_number==0 and b.locator_kind=='structure' and b.source_ref for b in doc.blocks)
        assert FACT in [b.text for b in doc.blocks]
        assert len([b for b in doc.blocks if b.block_type=='list_item'])==2
        table=next(b for b in doc.blocks if b.block_type=='table')
        assert '| Пожар' in table.text and 'Умысел' in table.text
        assert table.table_data['num_cols']==3 and table.table_data['num_rows']==2
        assert table.section=='Исключения'
        from rag_arbiter.application.exhaustive import ExhaustiveContextBatcher
        batcher=ExhaustiveContextBatcher(config.llm,QUESTION)
        batches=batcher.build([doc.model_dump()])
        assert batches and all(b.page_ranges==[] for b in batches)
        assert FACT in '\n'.join(u.text for batch in batches for u in batch.units)
        assert '"page":null' in batcher.request(batches[0].units).question.replace(' ','')
        assert [b.reading_order for b in doc.blocks]==list(range(len(doc.blocks)))
        repeated,again=Ingestor(config,store,NoOCR()).ingest([path])
        assert again['recognition_reused']==1 and repeated[0].document_id==doc.document_id
    finally:store.close()


@pytest.mark.parametrize('kind',['empty','broken'])
def test_empty_corrupt_docx_isolated(config,kind):
    bad=config.corpus_path/'bad.docx'
    if kind=='empty':WordDocument().save(bad)
    else:bad.write_bytes(b'not a zip')
    good=config.corpus_path/'good.docx';make_docx(good)
    store=MetadataStore(config.sqlite_path)
    try:
        docs,stats=Ingestor(config,store,NoOCR()).ingest([bad,good])
        assert len(docs)==1 and docs[0].file_name=='good.docx'
        assert len(stats['errors'])==1
    finally:store.close()


class GroundedFixtureLLM:
    def generate(self,request):
        if request.question.startswith('Переформулируй исходный вопрос'):
            return LLMResult(model='MOCK',status='SUCCESS',text=json.dumps({'retrieval_query':QUESTION},ensure_ascii=False))
        if request.context_type=='grounded_rag':
            parts=re.split(r'(?m)^(\{"reference":[^\n]+\})\n',request.context)
            ref=next((json.loads(parts[i])['reference'].strip('[]') for i in range(1,len(parts),2) if FACT in parts[i+1]),None)
            body=dict(answer=FACT,claims=[dict(claim_id='C1',text=FACT,supporting_source_ids=[ref])]) if ref else dict(answer='',claims=[],insufficient_context=True)
            return LLMResult(model='MOCK',status='SUCCESS',text=json.dumps(body,ensure_ascii=False))
        return LLMResult(model='MOCK',status='SUCCESS',text='')


def test_docx_mcp_search_chat_citations_versions(config,tmp_path):
    root=tmp_path/'Сетевая папка';root.mkdir()
    path=root/'Правила.docx';make_docx(path)
    config.network_sources.allowed_roots=[str(root)]
    config.network_sources.stability_seconds=0
    config.fixed_tokens=config.fixed_max_tokens=config.structure_max_tokens=1000
    config.fixed_overlap=100
    llm=GroundedFixtureLLM()
    with TestClient(create_app(config,factory(FakeRuntime()),lambda _:llm),base_url='http://127.0.0.1',headers=HEADERS) as client:
        sync=client.app.state.network
        source=sync.add({'name':'Local DOCX fixture','root_path':str(root)})
        first=scan(sync,source)
        assert first['ingestion_completed']==1,first
        file=sync.page('network_files',source['id'])[0]
        assert file['format']=='DOCX' and file['indexing_status']=='SUCCESS' and file['chunk_count']>0
        rid=file['run_id'];did=file['document_id']
        results=client.app.state.views.search(rid,QUESTION)
        hits=results['structure']['hits']
        assert any(FACT in h['text'] for h in hits)
        assert all(h['page_start']==0 and not h['source_pages'] for h in hits)
        from rag_arbiter.reranking import RerankedCandidate
        class Reranker:
            def status(self):return {'status':'READY','loaded':True,'model':'MOCK','device':'cpu'}
            def rerank(self,q,candidates):
                return [RerankedCandidate(**c.model_dump(),rerank_rank=i+1,rerank_score=.99) for i,c in enumerate(candidates)]
        client.app.state.rag.reranker_factory=lambda _:Reranker()
        chat=client.app.state.chat
        memory=MemoryProvider()
        chat.memory=TaskMemoryUpdater(lambda _:memory,config.llm)
        chat.resolver=ContextualQueryBuilder(lambda _:memory,config.llm)
        install(chat,llm,config.llm)
        sid=client.post('/api/chat/sessions',json={'processing_run_id':rid}).json()['session_id']
        response=client.post(f'/api/chat/sessions/{sid}/messages',json={'content':'/help '+QUESTION,'expected_version':0})
        assert response.status_code==202,response.text
        chat.futures[response.json()['turn_id']].result(timeout=40)
        turn=chat.load(sid)['turns'][0]
        assert turn['grounding_status']=='GROUNDED',turn
        assert turn['citations_json']
        with sync.runs.db() as store:
            for raw in turn['citations_json']:
                assert raw['locator_kind']=='structure' and raw['page_start']==0
                assert CitationBuilder(store).exact(Citation(**raw))
        c=turn['citations_json'][0]
        url=f"/ui/runs/{rid}/documents/{did}/blocks/{c['locator']['block_id']}"
        fragment=client.get(url)
        assert fragment.status_code==200 and FACT in fragment.text
        assert client.get(url.replace(did,'forbidden')).status_code==404
        html=client.get('/chat?session_id='+sid).text
        assert 'Открыть структурный фрагмент документа' in html and 'стр. 0' not in html
        repeat=scan(sync,source)
        assert repeat['unchanged_count']==1 and repeat['ingestion_queued']==0
        make_docx(path,'Новое условие: выплата производится в течение пяти дней.')
        changed=scan(sync,source)
        assert changed['modified_count']==changed['ingestion_completed']==1,changed
        latest=sync.page('network_files',source['id'])[0]
        assert latest['document_id']!=did
        snapshot=client.app.state.views.snapshot(latest['run_id'])
        assert all(did not in i['document_ids'] for i in snapshot['indexes'].values())
        # Explicitly labelled fixture evidence; contains no user documents or credentials.
        evidence=dict(classification='MOCK TEST',real_components=['Docling','MCP stdio','canonical SQLite','Qdrant','Chat orchestration','grounding'],
            mocked_components=['embeddings','reranker','generation','dialogue classifiers'],source_id=source['id'],scan_id=first['scan_id'],
            document_id=did,format='DOCX',ingestion_status=file['ingestion_status'],chunk_count=file['chunk_count'],index_status=file['indexing_status'],
            question=QUESTION,answer=chat.load(sid)['messages'][-1]['content'],citations=turn['citations_json'],retrieval=[{'document_id':h['document_id'],'chunk_id':h['chunk_id'],'score':h['score'],'text':h['text']} for h in hits],
            repeat_scan=repeat,modified_scan=changed)
        output=Path('data/day31-e2e');output.mkdir(exist_ok=True)
        (output/'mock-evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')

def test_local_root_dialog_approval_and_remote_denial(config,tmp_path):
    root=tmp_path/'operator-source';root.mkdir()
    form={'name':'Operator source','root_path':str(root),'recursive':'on','enabled':'on',
        'scan_interval':'300','include_patterns':'*','exclude_patterns':'','authorize_root':'yes'}
    with TestClient(create_app(config,factory(FakeRuntime())),base_url='http://127.0.0.1',headers=HEADERS,client=('192.0.2.1',50000)) as remote:
        assert remote.post('/api/sources/check-path',json={'root_path':str(root)}).status_code==403
        assert remote.post('/ui/sources',data=form).status_code==403
    with TestClient(create_app(config,factory(FakeRuntime())),base_url='http://127.0.0.1',headers=HEADERS,client=('127.0.0.1',50000)) as local:
        html=local.get('/sources?connect=1',headers={'HX-Request':'true'}).text
        assert '<dialog id="source-dialog">' in html and 'name="root_path"' in html
        assert local.post('/api/sources/check-path',json={'root_path':str(root)}).json()['status']=='ACCESS_OK'
        missing = local.post('/api/sources/check-path', json={'root_path': str(root/'missing')})
        assert missing.status_code == 400
        assert 'UNC' in missing.json()['detail']
        denied_form = local.post('/ui/sources', data={**form, 'root_path': str(root/'missing')})
        assert denied_form.status_code == 400

        assert local.post('/ui/sources',data={**form,'authorize_root':''}).status_code==400
        assert local.post('/ui/sources',data=form).status_code==200
        with local.app.state.runs.db() as store:
            assert store.all('network_root_grants')[0]['root_path']==str(root)
    config.network_sources.allowed_roots=[]
    with TestClient(create_app(config,factory(FakeRuntime())),base_url='http://127.0.0.1',headers=HEADERS) as restarted:
        assert str(root) in restarted.app.state.network.config.network_sources.allowed_roots
