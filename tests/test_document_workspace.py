import io,zipfile,time,threading
from contextlib import contextmanager
from types import SimpleNamespace
import pytest
from fastapi import HTTPException
from rag_arbiter.config import Config
from rag_arbiter.storage import MetadataStore
from rag_arbiter.application.auth import AuthService
from rag_arbiter.application.document_workspace import WorkspaceService,ToolInput,sha
from rag_arbiter.application.archive_publisher import ArchivePublisher
from rag_arbiter.application import workspace_formats as fmt

@pytest.fixture
def ws(tmp_path):
    config=Config(web_data_path=tmp_path/'web',sqlite_path=tmp_path/'db.sqlite')
    config.telegram_auth.admin_ids=[99901]
    class Runs:
        def __init__(self):self.config=config;self.mutation_lock=threading.RLock()
        @contextmanager
        def db(self):
            store=MetadataStore(config.sqlite_path)
            try:yield store
            finally:store.db.close()
    runs=Runs();auth=AuthService(runs);admin=auth.register(99901);client=auth.register(99902);auth.change(admin,client['id'],'approve');client=auth.user(99902)
    knowledge=SimpleNamespace(grants=lambda:{})
    service=WorkspaceService(runs,auth,knowledge)
    root=tmp_path/'archive';root.mkdir();config.workspace.publication_roots=[str(root)];config.network_sources.allowed_roots=[str(root)]
    with runs.db() as db:db.put('network_sources','source',dict(id='source',name='test',root_path=str(root),enabled=True))
    return service,admin,client,root

def draft(ws,user=None,format='docx'):
    service,admin,client,root=ws;user=user or admin
    task=service.task(user,'test')
    file=service.call(user,task['id'],ToolInput(operation='create_document',format=format,content='Проверяемый текст 123')).data
    return task,file

def proposal(ws,name='result.docx',operation='CREATE_NEW'):
    service,admin,_,root=ws;task,file=draft(ws)
    request=ToolInput(operation='prepare_publication',file_id=file['id'],source_id='source',target_path=name,publication_operation=operation)
    return task,file,ArchivePublisher(service).prepare(admin,request)

@pytest.mark.parametrize('format',['docx','pdf','txt','md'])
def test_real_file_roundtrip(ws,format):
    service,admin,_,_=ws;task,file=draft(ws,format=format)
    row,raw=service.file_bytes(admin,file['id'])
    assert any('123' in b['text'] for b in fmt.blocks(raw,format)['blocks'])
    assert service.validate(admin,file['id'])['structural_status']=='PASS'
    assert row['status']=='UNVERIFIED'

def test_edit_preserves_other_package_parts_and_previous_revision(ws):
    from docx import Document
    service,admin,_,_=ws;task=service.task(admin,'edit')
    document=Document();document.add_paragraph('Target');document.add_paragraph('Untouched',style='List Number');document.add_table(rows=1,cols=1).cell(0,0).text='Table preserved';document.sections[0].header.paragraphs[0].text='Header';out=io.BytesIO();document.save(out)
    first=service.save(admin,task,'docx','test',out.getvalue());raw=out.getvalue()
    second=service.call(admin,task['id'],ToolInput(operation='edit_document',file_id=first['id'],replacements={'Target':'Changed'},append='New section')).data
    _,new=service.file_bytes(admin,second['id']);assert service.file_bytes(admin,first['id'])[1]==raw
    with zipfile.ZipFile(io.BytesIO(raw)) as a,zipfile.ZipFile(io.BytesIO(new)) as b:
        assert all(a.read(n)==b.read(n) for n in a.namelist() if n!='word/document.xml')
    texts=[b['text'] for b in fmt.blocks(new,'docx')['blocks']]
    assert 'Changed' in texts and 'Untouched' in texts and 'Table preserved' in texts and 'New section' in texts
    diff=service.get('workspace_diffs',second['id'],admin)
    assert {c['change_type'] for c in diff['changes']}&{'ADDED','MODIFIED'}

def test_owner_and_client_publication(ws):
    service,admin,client,_=ws;task,file=draft(ws)
    with pytest.raises(HTTPException):service.file_bytes(client,file['id'])
    ct,cf=draft(ws,user=client)
    with pytest.raises(HTTPException):ArchivePublisher(service).prepare(client,ToolInput(operation='prepare_publication',file_id=cf['id'],source_id='source',target_path='x.docx'))
    with pytest.raises(HTTPException):service.call(admin,task['id'],ToolInput(operation='confirm_publication'))

@pytest.mark.parametrize('name',['../outside.docx','C:\\outside.docx','\\\\host\\share\\out.docx','x:stream.docx','NUL.docx','trailing. /a.docx'])
def test_bad_paths(ws,name):
    service,admin,_,_=ws
    with pytest.raises(ValueError):ArchivePublisher(service).target(admin,'source',name)

def test_allowlist(ws):
    service,admin,_,_=ws;service.runs.config.workspace.publication_roots=[]
    with pytest.raises(ValueError,match='ROOT_NOT_ALLOWED'):ArchivePublisher(service).target(admin,'source','x.docx')

def test_create_publish_idempotency_and_explicit_limitations(ws):
    service,admin,_,root=ws;task,file,p=proposal(ws);publisher=ArchivePublisher(service)
    with pytest.raises(ValueError,match='LIMITATIONS'):publisher.confirm(admin,p['id'],'key')
    result=publisher.confirm(admin,p['id'],'key',True)
    assert result['status']=='PUBLISHED_PENDING_INDEX',result
    assert sha((root/'result.docx').read_bytes())==file['sha256']
    assert publisher.confirm(admin,p['id'],'key',True)['new_hash']==file['sha256']
    with pytest.raises(ValueError,match='CONSUMED'):publisher.confirm(admin,p['id'],'different',True)
    assert result['sync_error']=='SYNC_START_FAILED'

def test_create_conflict(ws):
    service,admin,_,root=ws;task,file,p=proposal(ws);(root/'result.docx').write_bytes(b'other writer')
    result=ArchivePublisher(service).confirm(admin,p['id'],'key',True)
    assert result['status']=='CONFLICT' and (root/'result.docx').read_bytes()==b'other writer'

def test_draft_changed_and_expired(ws,monkeypatch):
    service,admin,_,root=ws;task,file,p=proposal(ws)
    monkeypatch.setattr('rag_arbiter.application.archive_publisher.time.time',lambda:p['expires_at']+1)
    with pytest.raises(ValueError,match='EXPIRED'):ArchivePublisher(service).confirm(admin,p['id'],'key',True)
    monkeypatch.undo()
    (service.root/admin['id']/(file['id']+'.docx')).write_bytes(b'changed')
    with pytest.raises(ValueError,match='DRAFT_CHANGED'):ArchivePublisher(service).confirm(admin,p['id'],'key',True)

def test_replace_backup_and_conflict(ws):
    service,admin,_,root=ws;old=fmt.create('docx','Old','Original');(root/'result.docx').write_bytes(old)
    with service.runs.db() as db:
        db.put('network_files','nf',dict(id='nf',source_id='source',relative_path='result.docx',content_hash=sha(old),document_id='v-old'))
        db.put('document_versions','logical-old',dict(version_id='logical-old',source_id='source',relative_path='result.docx',content_hash=sha(old)))
    task,file,p=proposal(ws,operation='REPLACE_EXISTING');result=ArchivePublisher(service).confirm(admin,p['id'],'key',True)
    assert result['status']=='PUBLISHED_PENDING_INDEX',result
    assert (root/('.workspace-old-'+p['id']+'.bak')).read_bytes()==old
    assert sha((root/'result.docx').read_bytes())==file['sha256']

def test_reparse_escape(ws,tmp_path):
    import subprocess
    service,admin,_,root=ws;outside=tmp_path/'outside';outside.mkdir();junction=root/'escape'
    made=subprocess.run(['cmd','/c','mklink','/J',str(junction),str(outside)],capture_output=True)
    if made.returncode:pytest.skip('junction unavailable')
    try:
        with pytest.raises(ValueError):ArchivePublisher(service).target(admin,'source','escape/result.docx')
    finally:junction.rmdir()

def test_revocation_and_cancel(ws):
    service,admin,client,_=ws;task,file=draft(ws,user=client)
    service.cancel(client,task['id'])
    with pytest.raises(ValueError,match='CANCELLED'):service.call(client,task['id'],ToolInput(operation='search_documents'))
    service.auth.change(admin,client['id'],'block')
    with pytest.raises(HTTPException):service.file_bytes(client,file['id'])

def test_failed_replace_restores_original(ws,monkeypatch):
    import rag_arbiter.application.archive_publisher as publisher
    service,admin,_,root=ws;old=fmt.create('docx','Old','Original');(root/'result.docx').write_bytes(old)
    with service.runs.db() as db:
        db.put('network_files','nf',dict(id='nf',source_id='source',relative_path='result.docx',content_hash=sha(old),document_id='v-old'))
        db.put('document_versions','logical-old',dict(version_id='logical-old',source_id='source',relative_path='result.docx',content_hash=sha(old)))
    _,_,p=proposal(ws,operation='REPLACE_EXISTING');original=publisher.rename_handle;calls=[]
    def fail_second(h,target):
        calls.append(str(target))
        if len(calls)==2:raise OSError('synthetic write failure')
        return original(h,target)
    monkeypatch.setattr(publisher,'rename_handle',fail_second)
    result=publisher.ArchivePublisher(service).confirm(admin,p['id'],'key',True)
    assert result['status']=='FAILED' and (root/'result.docx').read_bytes()==old

def test_changed_replace_target_conflicts(ws):
    service,admin,_,root=ws;old=fmt.create('docx','Old','Original');(root/'result.docx').write_bytes(old)
    with service.runs.db() as db:
        db.put('network_files','nf',dict(id='nf',source_id='source',relative_path='result.docx',content_hash=sha(old),document_id='v-old'))
        db.put('document_versions','logical-old',dict(version_id='logical-old',source_id='source',relative_path='result.docx',content_hash=sha(old)))
    _,_,p=proposal(ws,operation='REPLACE_EXISTING');changed=fmt.create('docx','Changed','External writer');(root/'result.docx').write_bytes(changed)
    result=ArchivePublisher(service).confirm(admin,p['id'],'key',True)
    assert result['status']=='CONFLICT' and (root/'result.docx').read_bytes()==changed

def test_proposal_immutable_and_foreign_confirmation(ws):
    import sqlite3
    service,admin,client,root=ws;_,_,p=proposal(ws)
    with service.runs.db() as db:
        with pytest.raises(sqlite3.IntegrityError):db.db.execute('UPDATE workspace_proposals SET data=? WHERE id=?',('{}',p['id']))
    with pytest.raises(HTTPException):ArchivePublisher(service).confirm(client,p['id'],'key',True)

def test_index_status_requires_exact_hash(ws):
    service,admin,_,root=ws;_,file,p=proposal(ws);publisher=ArchivePublisher(service);publisher.confirm(admin,p['id'],'key',True)
    with service.runs.db() as db:db.put('network_files','nf',dict(id='nf',source_id='source',relative_path='result.docx',content_hash='different',ingestion_status='SUCCESS',indexing_status='SUCCESS'))
    assert publisher.status(admin,p['id'])['status']=='PUBLISHED_PENDING_INDEX'
    with service.runs.db() as db:db.put('network_files','nf',dict(id='nf',source_id='source',relative_path='result.docx',content_hash=file['sha256'],ingestion_status='FAILED',indexing_status='FAILED'))
    assert publisher.status(admin,p['id'])['status']=='INDEX_FAILED'

def test_read_pins_version_and_acl(ws):
    service,admin,_,_=ws
    service.knowledge.grants=lambda:{'doc':dict(recognition_version='v1',title='test')}
    with service.runs.db() as db:db.put('normalized_documents','v1',dict(document_id='doc',content_hash='source-hash',blocks=[dict(text='Evidence')]))
    task=service.task(admin,'read');result=service.call(admin,task['id'],ToolInput(operation='read_document',document_id='doc',version_id='v1'))
    assert result.data['source']['version_id']=='v1'
    with pytest.raises(HTTPException):service.read(admin,'doc','v2')
    file=service.call(admin,task['id'],ToolInput(operation='create_document',content='draft')).data
    service.knowledge.grants=lambda:{}
    with pytest.raises(HTTPException):service.file_bytes(admin,file['id'])

@pytest.mark.parametrize('question',[
 'Найди субагентские договоры, сравни условия отчётности и подготовь документ DOCX',
 'Возьми подготовленное заключение, добавь раздел с различиями и сохрани новую редакцию',
 'Сохраните результат в PDF',
])
def test_explicit_file_goal_precedes_analysis(question):
    from rag_arbiter.application.document_workspace import explicit_file_goal
    assert explicit_file_goal(question)

@pytest.mark.parametrize('question',['Какие документы есть?','Кто создал документ?','Какие обязанности по договору?'])
def test_inventory_is_not_a_file_command(question):
    from rag_arbiter.application.document_workspace import explicit_file_goal
    assert not explicit_file_goal(question)


def test_unc_share_root_handle_spelling():
    from pathlib import PureWindowsPath
    assert PureWindowsPath(r"\\server\share") == PureWindowsPath("\\\\server\\share\\")
    assert PureWindowsPath(r"\\server\other") != PureWindowsPath(r"\\server\share")

@pytest.mark.parametrize('format',['txt','md'])
def test_nonindexable_publication_rejected(ws,format):
    service,admin,_,root=ws
    task=service.task(admin,'draft')
    file=service.call(admin,task['id'],ToolInput(operation='create_document',format=format,content='draft')).data
    with pytest.raises(ValueError,match='PUBLICATION_FORMAT_NOT_INDEXABLE'):
        ArchivePublisher(service).prepare(admin,ToolInput(operation='prepare_publication',file_id=file['id'],source_id='source',target_path='result.'+format))


def test_crash_after_create_reconciles_without_rewrite(ws):
    service,admin,_,root=ws;_,file,p=proposal(ws);publisher=ArchivePublisher(service)
    publisher.confirm(admin,p['id'],'key',True)
    row=service.get('workspace_publications',p['id'],admin);row['status']='PUBLISHING';service.put('workspace_publications',row)
    before=(root/'result.docx').stat().st_mtime_ns
    result=publisher.confirm(admin,p['id'],'key',True)
    assert result['recovery']=='COMPLETED_WRITE_RECONCILED'
    assert (root/'result.docx').stat().st_mtime_ns==before


def test_crash_after_backup_restores_original(ws):
    service,admin,_,root=ws;old=fmt.create('docx','Old','Original');target=root/'result.docx';target.write_bytes(old)
    with service.runs.db() as db:
        db.put('network_files','nf',dict(id='nf',source_id='source',relative_path='result.docx',content_hash=sha(old)))
        db.put('document_versions','logical-old',dict(version_id='logical-old',source_id='source',relative_path='result.docx',content_hash=sha(old)))
    _,file,p=proposal(ws,operation='REPLACE_EXISTING');assert p['previous_version']=='logical-old'
    target.rename(root/('.workspace-old-'+p['id']+'.bak'))
    service.put('workspace_publications',dict(**p,idempotency_key='key',status='PUBLISHING',new_hash=file['sha256'],accepted_limitations=True))
    result=ArchivePublisher(service).confirm(admin,p['id'],'key',True)
    assert result['recovery']=='ORIGINAL_RESTORED' and target.read_bytes()==old


def test_replace_requires_logical_version(ws):
    service,admin,_,root=ws;old=fmt.create('docx','Old','Original');(root/'result.docx').write_bytes(old)
    with service.runs.db() as db:db.put('network_files','nf',dict(id='nf',source_id='source',relative_path='result.docx',content_hash=sha(old),document_id='canonical-only'))
    with pytest.raises(ValueError,match='LOGICAL_VERSION_NOT_AVAILABLE'):proposal(ws,operation='REPLACE_EXISTING')
