import json,time,io,zipfile
import pytest
from pydantic import ValidationError
from test_document_workspace import ws,draft,proposal
from rag_arbiter.application.document_workspace import ToolInput
from rag_arbiter.application import workspace_formats as fmt
from rag_arbiter.llm import LLMResult

@pytest.mark.parametrize('data',[
 {'operation':'execute_shell'}, {'operation':'read_document'},
 {'operation':'edit_document','file_id':'x'}, {'operation':'compare_documents','file_id':'x'},
 {'operation':'prepare_publication','file_id':'x'}, {'operation':'create_document','unknown':True}])
def test_typed_invalid_calls(data):
    with pytest.raises(ValidationError):ToolInput(**data)

class Script:
    def __init__(self,values):self.values=iter(values);self.calls=[]
    def generate(self,req):
        self.calls.append(req)
        value=next(self.values)
        return LLMResult(model='fixture',status='SUCCESS',text=json.dumps(value) if isinstance(value,dict) else value)

def run(ws,values):
    service,admin,_,_=ws;provider=Script(values)
    result=service.execute_goal(admin,'Create a document',lambda cfg:provider,service.runs.config.llm)
    return result,provider

def test_invalid_json_one_repair_then_stops(ws):
    result,p=run(ws,['invalid','invalid'])
    assert result['status']=='PARTIAL' and not result['files'] and result['repair_count']==1 and len(p.calls)==2
    assert result['error_reason']=='INVALID_TOOL_CALL'

def test_backend_completion_needs_no_extra_planner_call(ws):
    result,p=run(ws,[dict(operation='create_document',content='Once')])
    assert len(result['files'])==1 and result['error_reason'] is None and len(p.calls)==1

def test_truncated_response_bounded(ws):
    service,admin,_,_=ws;calls=[]
    class Truncated:
        def generate(self,req):calls.append(1);return LLMResult(model='fixture',status='TRUNCATED',text='{')
    result=service.execute_goal(admin,'Create',lambda cfg:Truncated(),service.runs.config.llm)
    assert len(calls)==1 and result['status']=='PARTIAL' and result['error_reason']=='OUTPUT_TRUNCATED'

def test_oversized_prompt_stops_before_provider(ws):
    service,admin,_,_=ws;cfg=service.runs.config.llm.model_copy(deep=True);cfg.full_document_context_budget=4097
    p=Script([]);result=service.execute_goal(admin,'Create',lambda cfg:p,cfg)
    assert not p.calls and result['status']=='PARTIAL'

def test_cancel_during_provider_prevents_mutation(ws):
    service,admin,_,_=ws
    class Cancel:
        def generate(self,req):
            with service.runs.db() as db:key=db.all('workspace_tasks')[-1]['id']
            service.cancel(admin,key)
            return LLMResult(model='fixture',status='SUCCESS',text=json.dumps(dict(operation='create_document',content='Must not exist')))
    result=service.execute_goal(admin,'Create',lambda cfg:Cancel(),service.runs.config.llm)
    assert result['status']=='CANCELLED' and not result['files']

def test_timeout_after_provider_prevents_mutation_and_releases(ws):
    service,admin,_,_=ws;service.cfg.timeout=.2
    class Slow:
        def generate(self,req):
            time.sleep(.3)
            return LLMResult(model='fixture',status='SUCCESS',text=json.dumps(dict(operation='create_document',content='Late')))
    result=service.execute_goal(admin,'Create',lambda cfg:Slow(),service.runs.config.llm)
    assert result['status']=='PARTIAL' and result['error_reason']=='WORKSPACE_TIMEOUT' and not result['files']
    service.cfg.timeout=30
    result,_=run(ws,[dict(operation='create_document',content='Next'),dict(operation="finish_task")])
    assert len(result['files'])==1

def test_compact_state_omits_full_trace(ws):
    service,admin,_,_=ws;task=service.task(admin,'compact')
    task['tools']=[dict(input=dict(operation='read_document'),output=dict(data=dict(workspace_text='x'*100000,blocks=['secret trace']*1000,citations=['large']*1000))) for _ in range(10)]
    payload=json.dumps(service.planner_state(task,[]))
    assert len(payload)<40000 and 'secret trace' not in payload

def test_docx_structured_creation_and_edit():
    from docx import Document
    raw=fmt.create('docx','Title','## Heading\n- Bullet\n| A | B |\n|---|---|\n| 1 | 2 |\n**Bold** text')
    doc=Document(io.BytesIO(raw));assert doc.tables[0].cell(1,1).text=='2'
    assert any(p.style.name=='Heading 2' for p in doc.paragraphs)
    assert any(p.style.name=='List Bullet' for p in doc.paragraphs)
    assert any(r.bold for p in doc.paragraphs for r in p.runs)
    changed=fmt.edit(raw,'docx',{},'## Added\n- Another')
    with zipfile.ZipFile(io.BytesIO(raw)) as a,zipfile.ZipFile(io.BytesIO(changed)) as b:
        assert all(a.read(n)==b.read(n) for n in a.namelist() if n!='word/document.xml')

def test_semantic_failure_is_not_structure_failure(ws):
    service,admin,_,_=ws;_,file=draft(ws)
    validation=service.validate_content(admin,file['id'],lambda cfg:Script([]),service.runs.config.llm)
    assert validation['structural_status']=='PASS' and validation['semantic_status']=='UNVERIFIED'
    assert validation['claims']==[] and validation['failure']['correlation_id']

def test_safe_exception_trace_excludes_payload(ws):
    from rag_arbiter.application.failure_diagnostics import capture
    service,_,_,_=ws
    try:raise RuntimeError('secret token payload')
    except Exception as exc:result=capture(service.runs,exc,route='/test')
    assert result['traceback'][0]['frames'] and 'secret token payload' not in json.dumps(result)


def test_recorded_direct_typed_tool_response():
    from rag_arbiter.application.document_workspace import FlatToolInput
    wire='{"operation":"read_document","document_id":"canonical-id","version_id":"pinned-version"}'
    assert FlatToolInput.model_validate_json(wire).root.operation=='read_document'
    with pytest.raises(ValidationError):FlatToolInput.model_validate_json('{"operation":"read_document","document_id":"x","version_id":"v","done":true}')

def test_resume_preserves_completed_mutations(ws):
    service,admin,_,_=ws
    first,_=run(ws,[dict(operation='create_document',content='Saved')])
    provider=Script([dict(operation="finish_task")])
    resumed=service.execute_goal(admin,'Create a document',lambda cfg:provider,service.runs.config.llm,resume_task_id=first['id'])
    assert resumed['id']==first['id'] and resumed['files']==first['files'] and len(resumed['tools'])==1
    assert resumed['attempt_count']==2 and resumed['previous_failures']


def test_recorded_redundant_tool_call_must_match_exactly():
    from rag_arbiter.application.document_workspace import FlatToolInput
    wire=dict(operation='read_document',document_id='d',version_id='v')
    with pytest.raises(ValidationError):FlatToolInput.model_validate({**wire,'tool':wire})
    with pytest.raises(ValidationError):FlatToolInput.model_validate({**wire,'tool':{**wire,'document_id':'other'}})

def test_planner_has_one_schema(ws):
    result,p=run(ws,[dict(operation="finish_task")])
    payload=json.loads(p.calls[0].question.split('\n',1)[1])
    assert 'schema' in payload and 'tools' not in payload


def test_false_completion_without_file(ws):
    result,p=run(ws,[dict(operation="finish_task")])
    assert result['status']=='PARTIAL' and result['error_reason']=='COMPLETION_WITHOUT_FILE'
    assert not result['files']


def test_completed_artifact_has_saved_validation(ws):
    service,admin,_,_=ws
    result,p=run(ws,[dict(operation='create_document',content='A real draft'),dict(operation="finish_task")])
    assert result['files']==result['saved_files']
    with service.runs.db() as db:checks=db.all('workspace_validations')
    assert all(any(v['file_id']==key and v['structural_status']=='PASS' for v in checks) for key in result['files'])


@pytest.mark.parametrize('value',[{},dict(done=True,tool=dict(operation='search_documents')),dict(operation='read_document'),dict(operation='unknown'),dict(done=True,extra=1)])
def test_canonical_contract_rejects_incomplete_or_conflicting(value):
    from rag_arbiter.application.document_workspace import FlatToolInput
    with pytest.raises(ValidationError):FlatToolInput.model_validate(value)



def test_repaired_schema_clears_active_error_and_keeps_trace(ws):
    result,p=run(ws,['invalid',dict(operation='create_document',content='Saved after repair'),dict(operation="finish_task")])
    assert result['error_reason'] is None and result['repair_count']==1
    assert result['planner_responses'][0]['text']=='invalid' and len(result['files'])==1
