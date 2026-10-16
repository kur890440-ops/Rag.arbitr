"""Source isolation through real MCP/Qdrant; inference is explicitly scripted."""
import json
import time
import pytest
from pathlib import Path
from docx import Document
from rag_arbiter.llm import LLMResult
from rag_arbiter.application.chat_memory import TaskMemoryUpdater, ContextualQueryBuilder
from rag_arbiter.application.chat_models import TaskState, MemoryItem
from rag_arbiter.application.retrieval_policy import RetrievalPolicyBuilder, ConstraintPredicate
from rag_arbiter.application.views import ResultsService
from rag_arbiter.application.exhaustive import ExhaustiveNoRAG
from test_rag import ragweb
from test_chat import MemoryProvider
from test_network_sources import scan
from dialogue_fixtures import install

FACTS=['Ущерб от наводнения покрывается страховкой.', 'Выплата производится в течение пяти дней.', 'Оплата через десять дней.']


def document(path,fact):
    doc=Document();doc.add_heading('Условия',1);doc.add_paragraph(fact);doc.save(path)


def test_source_chat_isolation_discovery_persistence_and_modes(ragweb,tmp_path):
    client,manual_run,llm,_=ragweb
    network=client.app.state.network;chat=client.app.state.chat;runs=client.app.state.runs
    # Keep real orchestration, storage, MCP, Docling and Qdrant. Only inference is scripted.
    original=llm.generate
    def generate(request):
        if request.question.startswith('Переформулируй исходный вопрос'):
            return LLMResult(model='MOCK',status='SUCCESS',text=json.dumps({'retrieval_query':'Какие условия предусмотрены?'}))
        if request.context_type=='grounded_rag':
            import re
            parts=re.split(r'(?m)^(\{"reference":[^\n]+\})\n',request.context)
            matches=[(fact,json.loads(parts[i])['reference'].strip('[]')) for i in range(1,len(parts),2) for fact in FACTS if fact in parts[i+1]]
            fact,ref=matches[0] if matches else ('',None)
            return LLMResult(model='MOCK',status='SUCCESS',text=json.dumps(dict(answer=fact,claims=[dict(claim_id='C1',text=fact,supporting_source_ids=[ref])] if ref else [],insufficient_context=not bool(ref))))
        return original(request)
    llm.generate=generate
    memory=MemoryProvider();chat.memory=TaskMemoryUpdater(lambda _:memory,runs.config.llm);chat.resolver=ContextualQueryBuilder(lambda _:memory,runs.config.llm)
    install(chat,llm,runs.config.llm)
    sources=[];ids=[]
    for n in range(2):
        root=tmp_path/f'source{n}';root.mkdir();document(root/'fixture.docx',FACTS[n])
        network.config.network_sources.allowed_roots.append(str(root))
        network.config.network_sources.stability_seconds=0
        source=network.add(dict(name=f'SOURCE {n}',root_path=str(root)));sources.append(source)
        result=scan(network,source);assert result['ingestion_completed']==1,result
        ids.append(network.page('network_files',source['id'])[0]['document_id'])
    session=chat.create(manual_run);sid=session.session_id
    version=0
    def select(scope,source=None):
        nonlocal version
        r=client.patch(f'/api/chat/sessions/{sid}/source',json=dict(scope=scope,source_id=source,expected_version=version))
        assert r.status_code==200,r.text
        version=r.json()['scope_version']
    def ask(question='/help Какие условия предусмотрены?'):
        current=chat.load(sid)
        turn=chat.start(sid,question,current['task_state']['version'])
        chat.futures[turn.turn_id].result(timeout=45)
        result=chat.load(sid)['turns'][-1]
        assert result['status']=='COMPLETED',result
        return result
    for n in range(2):
        select('NETWORK',sources[n]['id'])
        turn=ask()
        assert turn['grounding_status']=='GROUNDED',turn
        assert turn['trace']['eligible_document_ids']==[ids[n]]
        assert {s['document_id'] for s in turn['sources_json']}=={ids[n]}
        assert ids[1-n] not in json.dumps(turn['trace']['retrieved_candidates'])
        assert FACTS[1-n] not in turn['response']['answer_text']
        follow=ask('Какие условия предусмотрены?')
        assert follow['trace']['source_selection']['source_id']==sources[n]['id']
        assert {s['document_id'] for s in follow['sources_json']}=={ids[n]}
        whole=ask('/help Сравни все условия в документах.')
        assert whole['trace']['answer_mode']=='EXHAUSTIVE_NO_RAG',whole
        assert whole['trace']['coverage']['eligible_document_ids']==[ids[n]]
        assert all(e['document_id']==ids[n] for f in whole['trace']['coverage']['merged_findings'] for e in f['evidence'])
        html=client.get('/chat',params={'session_id':sid}).text
        assert 'Выбрать документы' in html and 'id="chat-discovery"' not in html
        assert f'/documents/{ids[n]}/blocks/' in html
        assert f'/ui/runs/{turn["trace"]["document_runs"][ids[n]]}/documents/' in html
        panel=client.get(f'/ui/chat/{sid}/discovery').text
        assert 'INDEXED' in panel and 'fixture.docx' in panel and 'Открыть документ' in panel
        assert client.get(f'/api/chat/sessions/{sid}').json()['session']['source_id']==sources[n]['id']
    select('ALL_DOCUMENTS')
    other=chat.create(manual_run)
    response=client.patch(f'/api/chat/sessions/{sid}/source',json=dict(scope='ALL_DOCUMENTS',expected_version=version,selected_document_ids=ids))
    assert response.status_code==200,response.text
    version=response.json()['scope_version']
    assert response.json()['selected_document_ids']==ids
    assert chat.load(other.session_id)['session']['selected_document_ids']==[]
    html=client.get('/chat',params={'session_id':sid}).text
    assert 'Выбрать документы' in html and 'href="/catalog"' not in html
    assert 'Анализ: 2 документов' in html
    selector=client.get(f'/api/chat/sessions/{sid}/documents')
    assert selector.status_code==200 and 'checkbox' in selector.text
    assert set(ids)<=set(client.get(f'/api/chat/sessions/{sid}/documents?all_ready=true').json()['document_ids'])
    multi=ask();assert multi['grounding_status']=='GROUNDED'
    assert set(multi['trace']['eligible_document_ids'])==set(ids)
    assert {s['document_id'] for s in multi['sources_json']}<=set(ids)
    assert {s['document_id'] for s in multi['trace']['retrieved_candidates']}<=set(ids)
    full=ask('/help Сравни все условия в документах.')
    assert set(full['trace']['coverage']['eligible_document_ids'])==set(ids)
    follow=ask('Какие условия предусмотрены?')
    assert set(follow['trace']['eligible_document_ids'])==set(ids)
    status=ask('/help');assert status['trace']['archive_status']['counts']['discovered']==2
    denied=client.patch(f'/api/chat/sessions/{sid}/source',json=dict(scope='NETWORK',source_id=sources[0]['id'],expected_version=version,selected_document_ids=ids))
    assert denied.status_code==400
    assert client.patch(f'/api/chat/sessions/{sid}/source',json=dict(scope='ALL_DOCUMENTS',expected_version=version,selected_document_ids=['not-indexed'])).status_code==400
    assert chat.load(sid)['session']['selected_document_ids']==ids
    select('ALL_DOCUMENTS');all_turn=ask()
    assert set(ids)<=set(all_turn['trace']['eligible_document_ids'])
    select('MANUAL');manual=ask()
    assert not set(ids)&set(manual['trace']['eligible_document_ids'])
    assert manual['trace']['eligible_document_ids']
    from rag_arbiter.application.chat import ChatSessionService
    restarted=ChatSessionService(runs,client.app.state.rag)
    assert restarted.load(sid)['session']['source_scope']=='MANUAL'
    with pytest.raises(ValueError,match='другой вкладке'):
        chat.start(sid,'Какие условия?',chat.load(sid)['task_state']['version'],expected_scope_version=version-1)
    # Unknown source and stale selector writes fail closed.
    assert client.patch(f'/api/chat/sessions/{sid}/source',json=dict(scope='NETWORK',source_id='missing',expected_version=version)).status_code==400
    assert client.patch(f'/api/chat/sessions/{sid}/source',json=dict(scope='ALL_DOCUMENTS',expected_version=version-1)).status_code==400
    select('NETWORK',sources[0]['id'])
    snap=ResultsService(runs).chat_snapshot(manual_run)
    with runs.db() as store:
        from rag_arbiter.application.document_metadata import enrich_documents
        metadata=enrich_documents(store,snap['indexes']['structure']['document_ids'],snap['corpus']['corpus_id'])
    state=TaskState(session_id=sid,constraints_json=[MemoryItem(key='exclude',value='exclude A',evidence='exclude A',source_message_id='user',predicates=[ConstraintPredicate(type='DOCUMENT_ID_EXCLUDE',values=[ids[0]])])])
    policy=RetrievalPolicyBuilder().build(state,'ALL_DOCUMENTS',snap['corpus']['corpus_id'],metadata,source_scope='NETWORK',source_id=sources[0]['id'],source_document_ids=[ids[0]])
    assert not policy.eligible_document_ids and policy.exclude_filters
    # Error remains visible with the file list; not converted into an empty list.
    with runs.db() as store:
        row=network.page('network_scans',sources[0]['id'],limit=1)[0]
        row.update(status='ERROR',errors=['MCP_TIMEOUT'],error_count=1)
        store.put('network_scans',row['scan_id'],row)
    panel=client.get(f'/ui/chat/{sid}/discovery').text
    assert 'ERROR' in panel and 'MCP_TIMEOUT' in panel and 'fixture.docx' in panel
    repeat=scan(network,sources[0]);assert repeat['unchanged_count']==1 and repeat['ingestion_queued']==0
    # A newly indexed file appears in the same saved session without choosing a new run.
    document(Path(sources[0]['root_path'])/'second.docx',FACTS[0])
    added=scan(network,sources[0]);assert added['new_count']==1 and added['unchanged_count']==1
    after=ask()
    current_ids={f['document_id'] for f in network.page('network_files',sources[0]['id'])}
    assert set(after['trace']['eligible_document_ids'])==current_ids
    assert 'second.docx' in client.get(f'/ui/chat/{sid}/discovery').text
    whole=ask('/help Сравни все условия в документах.')
    assert set(whole['trace']['coverage']['eligible_document_ids'])==current_ids
    # Active canonical manifest controls eligibility, not a stale SUCCESS label alone.
    with runs.db() as store:
        for row in network.page('network_files',sources[0]['id']):
            row['active_run_id']=None;store.put('network_files',row['id'],row)
    empty=ask();assert empty['response']['status']=='NO_INDEXED_DOCUMENTS'
    assert not empty['sources_json']
    assert 'INDEXED</td>' not in client.get(f'/ui/chat/{sid}/discovery').text


def test_automatic_scan_wiring_and_new_file(ragweb,tmp_path):
    client,_,_,_=ragweb;svc=client.app.state.network
    root=tmp_path/'automatic';root.mkdir()
    svc.config.network_sources.allowed_roots.append(str(root));svc.config.network_sources.stability_seconds=0
    source=svc.add(dict(name='Automatic fixture',root_path=str(root),scan_interval=1))
    svc.config.network_sources.enabled=True
    try:
        deadline=time.monotonic()+20
        while not svc.page('network_scans',source['id']) and time.monotonic()<deadline:time.sleep(.1)
        assert svc.scheduler_running
        document(root/'automatic.docx',FACTS[0])
        deadline=time.monotonic()+45
        while time.monotonic()<deadline:
            rows=svc.page('network_files',source['id'])
            if rows and rows[0].get('indexing_status')=='SUCCESS':break
            time.sleep(.2)
        assert rows and rows[0]['status']=='NEW',rows
        info=svc.discovery(source['id'])
        assert info['automatic'] and info['counts']['indexed']==1
        assert info['active'] or info['next_scan']
    finally:svc.config.network_sources.enabled=False
